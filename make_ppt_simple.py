# -*- coding: utf-8 -*-
"""產生「功能介紹」簡報（4 頁，16:9）—— 給口頭報告用的精簡版。

完整版（9 頁，含數字與發現）請跑 make_ppt.py。
用法：python make_ppt_simple.py  →  outputs/專題功能介紹.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree

C = dict(gold=RGBColor(0xC8,0xA9,0x51), green=RGBColor(0x1E,0x8C,0x66),
         dark=RGBColor(0x26,0x30,0x2B), grey=RGBColor(0x7A,0x84,0x7E),
         light=RGBColor(0xF2,0xF3,0xF2), white=RGBColor(0xFF,0xFF,0xFF),
         mint=RGBColor(0xE2,0xF0,0xEA), line=RGBColor(0xE4,0xE7,0xE5))
FONT = "微軟正黑體"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333


def _ea(run, name):
    """python-pptx 只設 latin 字型，中文要另外補 <a:ea>。"""
    rPr = run._r.get_or_add_rPr()
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for tag in ("ea", "cs"):
        el = rPr.find(ns + tag)
        if el is None:
            el = etree.SubElement(rPr, ns + tag)
        el.set("typeface", name)


def text(slide, x, y, w, h, runs, size=18, color="dark", bold=False,
         align=PP_ALIGN.LEFT, spacing=1.3):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(runs if isinstance(runs, list) else [runs]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for st, opt in (line if isinstance(line, list) else [(line, {})]):
            r = p.add_run(); r.text = st
            f = r.font
            f.size = Pt(opt.get("size", size)); f.bold = opt.get("bold", bold)
            f.color.rgb = C[opt.get("color", color)]; f.name = FONT
            _ea(r, FONT)
    return tb


def card(slide, x, y, w, h, fill="white", edge="line"):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = C[fill]
    if edge:
        sh.line.color.rgb = C[edge]; sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def bar(slide, x, y, w, h, col):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = C[col]; sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def head(title, sub=None):
    s = prs.slides.add_slide(BLANK)
    bar(s, 0.9, 0.68, 0.055, 0.6, "gold")
    text(s, 1.15, 0.62, 11.5, 0.6, title, size=30, bold=True)
    if sub:
        text(s, 1.15, 1.35, 11, 0.4, sub, size=15, color="grey")
    return s


# ========== 1. 封面 ==========
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = C["dark"]; bg.line.fill.background()
bg.shadow.inherit = False
bar(s, 1.5, 2.7, 0.075, 0.8, "gold")
text(s, 1.5, 2.0, 10, 0.4, "專題功能介紹", size=16, color="gold")
text(s, 1.85, 2.62, 11, 1.0, "黃金避險決策系統", size=46, bold=True, color="white")
text(s, 1.85, 3.75, 10.5, 0.6,
     "幫持有黃金的人判斷：現在該不該避險、該避多少", size=18, color="light")
text(s, 1.5, 5.7, 11, 0.5,
     [[("網站已上線　", {"color": "gold"}),
       ("lilianlee1025-eng.github.io/GoldHedging", {"color": "white"})]], size=14)

# ========== 2. 網站功能 ==========
s = head("網站有哪些功能", "五個分頁，從「看現況」到「怎麼做」再到「有沒有效」")
feats = [("儀表板", "看金價走勢、模型預測的未來 5 天，以及下跌機率與避險建議", "gold"),
         ("避險策略模擬器", "輸入自己手上的黃金，試算期貨、保護性賣權、Collar 各自的損益", "green"),
         ("期貨 vs 期權原理", "拖動滑桿，即時看兩種避險工具的損益曲線差在哪", "gold"),
         ("回測結果", "把策略套回 2008～2026 的真實歷史，驗證避險到底有沒有用", "green"),
         ("模型說明", "LSTM 架構、六個輸入因子，以及三種模型的表現比較", "gold")]
for i, (t, d, col) in enumerate(feats):
    y = 2.15 + i * 1.0
    card(s, 0.9, y, 11.55, 0.86)
    bar(s, 0.9, y + 0.13, 0.05, 0.6, col)
    text(s, 1.25, y + 0.19, 3.1, 0.45, t, size=17, bold=True)
    text(s, 4.5, y + 0.22, 7.8, 0.45, d, size=13.5, color="grey")

# ========== 3. 個人化與會員登入 ==========
s = head("個人化與會員登入", "同一個模型，每個人看到的建議不一樣")
items = [("風險偏好", "保守／穩健／積極三種，決定下跌機率到多少才開始避險", "green"),
         ("多筆部位", "金條用台錢、存摺用公克、ETF 用 GLD 股、金幣用盎司，可以分開記", "gold"),
         ("記住你的設定", "計價單位、圖表類型等都會記住，重新整理不會歸零", "green")]
for i, (t, d, col) in enumerate(items):
    y = 2.15 + i * 1.0
    card(s, 0.9, y, 11.55, 0.86)
    bar(s, 0.9, y + 0.13, 0.05, 0.6, col)
    text(s, 1.25, y + 0.19, 2.9, 0.45, t, size=17, bold=True)
    text(s, 4.3, y + 0.22, 8.0, 0.45, d, size=13.5, color="grey")

card(s, 0.9, 5.3, 11.55, 1.6, fill="mint", edge=None)
bar(s, 0.9, 5.3, 0.05, 1.6, "green")
text(s, 1.3, 5.55, 10.9, 1.2, [
    [("會員登入", {"bold": True, "size": 18, "color": "green"})],
    [("登入後，上面這些設定會同步到雲端，換手機或換電腦都看得到；"
      "沒有登入也能完整使用，設定會存在自己的瀏覽器裡。", {"size": 14, "color": "dark"})],
], size=14)

# ========== 4. 資料庫怎麼架 ==========
s = head("資料庫怎麼架", "用 Supabase（PostgreSQL），為了支援多筆部位做了正規化")
tables = [("profiles", "風險偏好與\n避險試算參數", "一對一", "green"),
          ("positions", "使用者持有的\n黃金部位", "一對多", "gold"),
          ("ui_prefs", "介面顯示\n偏好設定", "一對一", "green")]
# auth.users（Supabase 內建）
card(s, 0.9, 2.5, 2.5, 1.5, fill="light", edge=None)
text(s, 1.15, 2.75, 2.0, 0.4, "auth.users", size=15, bold=True)
text(s, 1.15, 3.2, 2.05, 0.6, "Supabase 內建的\n帳號表".split("\n"),
     size=11.5, color="grey")
for i, (name, desc, rel, col) in enumerate(tables):
    x = 4.6 + i * 2.85
    card(s, x, 2.15, 2.5, 2.2)
    bar(s, x, 2.15, 2.5, 0.07, col)
    text(s, x + 0.25, 2.42, 2.0, 0.4, name, size=15, bold=True)
    text(s, x + 0.25, 2.9, 2.05, 0.8, desc.split("\n"), size=11.5, color="grey")
    text(s, x + 0.25, 3.85, 2.0, 0.3, rel, size=11, color=col, bold=True)
    text(s, 3.45, 2.9 + i * 0.0, 1.1, 0.3, "", size=10)
text(s, 3.45, 2.95, 1.1, 0.4, "1 : 1\n1 : N\n1 : 1".split("\n"),
     size=11, color="grey", align=PP_ALIGN.CENTER)

rows = [("為什麼要拆三張表",
         "原本設定全部塞在一個欄位裡。改成一位使用者可以有「多筆」部位後，"
         "就把部位獨立成 positions 表（一對多），偏好與顯示設定各自一張，資料不重複也好維護。"),
        ("怎麼確保只看得到自己的資料",
         "三張表都開啟 Row Level Security，四條政策（查詢／新增／修改／刪除）"
         "都規定 auth.uid() = user_id，所以每個人只能讀寫自己那一列。"),
        ("沒登入怎麼辦",
         "設定存在瀏覽器本機（localStorage）；登入後比對雲端與本機的更新時間，誰新用誰。")]
for i, (t, d) in enumerate(rows):
    y = 4.6 + i * 0.85
    text(s, 0.9, y, 3.5, 0.4, t, size=13.5, bold=True, color="green")
    text(s, 4.6, y, 7.9, 0.7, d, size=12, color="grey", spacing=1.2)

out = "outputs/專題功能介紹.pptx"
prs.save(out)
print("已產生：", out, f"（{len(prs.slides._sldIdLst)} 頁）")
